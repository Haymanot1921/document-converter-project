import sys
import tempfile
from pathlib import Path
import fitz
from pptx import Presentation
from pptx.util import Inches
from PIL import Image

def main(source_arg: str, destination_arg: str) -> None:
    source = Path(source_arg).resolve()
    destination = Path(destination_arg).resolve()
    if not source.is_file() or source.suffix.lower() != ".pdf":
        raise ValueError("Input must be an existing PDF file")
    destination.parent.mkdir(parents=True, exist_ok=True)

    pdf = fitz.open(source)
    if pdf.page_count < 1:
        raise ValueError("PDF contains no pages")
    if pdf.page_count > 100:
        raise ValueError("PDF has too many pages (maximum 100)")

    presentation = Presentation()
    # Widescreen 16:9 slide size
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    with tempfile.TemporaryDirectory(prefix="pdfppt-") as temp_dir:
        for index, page in enumerate(pdf):
            pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
            image_path = Path(temp_dir) / f"page-{index + 1}.png"
            pix.save(image_path)
            slide = presentation.slides.add_slide(blank)
            with Image.open(image_path) as image:
                img_w, img_h = image.size
            slide_ratio = presentation.slide_width / presentation.slide_height
            image_ratio = img_w / img_h
            if image_ratio > slide_ratio:
                width = presentation.slide_width
                height = int(width / image_ratio)
                left = 0
                top = int((presentation.slide_height - height) / 2)
            else:
                height = presentation.slide_height
                width = int(height * image_ratio)
                top = 0
                left = int((presentation.slide_width - width) / 2)
            slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)

    pdf.close()
    presentation.save(destination)
    if not destination.is_file() or destination.stat().st_size == 0:
        raise RuntimeError("PPTX output was not created")

if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: pdf_to_ppt.py input.pdf output.pptx")
    main(sys.argv[1], sys.argv[2])
